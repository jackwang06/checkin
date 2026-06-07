"""胶水层：把名单 xlsx 规范化为标准 roster.csv。

用法：uv run python scripts/xlsx_to_csv.py [--xlsx 晚点名线上表格.xlsx] [--out csv/roster.csv]

输入 xlsx 列：班级名称 | 学号 | 姓名 | (在校状态等列忽略——那是某时刻快照，不属于名单)
输出 csv 列：student_id,student_name,grade,major,class_no,class_tag,class_full_name
编码 UTF-8 带 BOM（Excel 双击打开不乱码）。

学号重复时保留首条，全部冲突行（含保留的那条）另存 csv/roster_conflicts.csv
供人工裁决；修正源 xlsx 后重跑本脚本 + import_csv.py 即可（全链路幂等）。
"""

import argparse
import csv
import re
import sys
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent

# '计算机23-1（卓越班）' -> 专业=计算机, 年级=23, 班号=1, 标签=卓越班
CLASS_RE = re.compile(r"^(.+?)(\d{2})-(\d+)(?:（(.*)）)?$")

CSV_HEADER = ["student_id", "student_name", "grade", "major",
              "class_no", "class_tag", "class_full_name"]


def parse_class_name(full_name: str) -> tuple[str, str, int, str]:
    m = CLASS_RE.match(full_name.strip())
    if not m:
        raise ValueError(f"班级名无法解析: {full_name!r}")
    major, grade, class_no, tag = m.groups()
    return major, grade, int(class_no), tag or ""


def convert(xlsx: Path, out: Path) -> Path:
    """xlsx 名单 → 标准 roster csv，返回输出路径。冲突另存 <out目录>/roster_conflicts.csv。"""
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb[wb.sheetnames[0]]

    raw = []                                     # (sid, name, class_full_name)
    for cls, sid, name, *_ in ws.iter_rows(min_row=2, values_only=True):
        if cls is None and sid is None:          # 末尾空行
            continue
        if not (cls and sid and name):
            print(f"警告: 跳过不完整行 {(cls, sid, name)}", file=sys.stderr)
            continue
        raw.append((str(sid).strip(), str(name).strip(), str(cls).strip()))

    id_count = {}
    for sid, _, _ in raw:
        id_count[sid] = id_count.get(sid, 0) + 1

    rows, conflicts, seen_ids = [], [], set()
    for sid, name, cls in raw:
        if id_count[sid] > 1:
            kept = sid not in seen_ids
            conflicts.append([sid, name, cls, "保留" if kept else "丢弃"])
            if not kept:
                print(f"警告: 学号重复 {sid}（{name}，{cls}），保留首条", file=sys.stderr)
        if sid in seen_ids:
            continue
        seen_ids.add(sid)
        major, grade, class_no, tag = parse_class_name(cls)
        rows.append([sid, name, grade, major, class_no, tag, cls])

    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(CSV_HEADER)
        w.writerows(rows)

    conflicts_path = out.parent / "roster_conflicts.csv"
    if conflicts:
        with open(conflicts_path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["student_id", "student_name", "class_full_name", "处理"])
            w.writerows(conflicts)
        print(f"⚠ 学号冲突 {len(conflicts)} 行已写入 {conflicts_path}，请人工裁决后修正源 xlsx",
              file=sys.stderr)

    classes = {r[6] for r in rows}
    majors = {r[3] for r in rows}
    grades = {r[2] for r in rows}
    print(f"已写入 {out}")
    print(f"学生 {len(rows)} 人 | 班级 {len(classes)} 个 | "
          f"专业 {len(majors)} 个 ({'/'.join(sorted(majors))}) | "
          f"年级 {len(grades)} 个 ({'/'.join(sorted(grades))})")
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsx", default=PROJECT_ROOT / "晚点名线上表格.xlsx", type=Path)
    ap.add_argument("--out", default=PROJECT_ROOT / "csv" / "roster.csv", type=Path)
    args = ap.parse_args()
    convert(args.xlsx, args.out)


if __name__ == "__main__":
    main()
